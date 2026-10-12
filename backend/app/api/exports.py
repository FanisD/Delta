# ruff: noqa: E501
from __future__ import annotations

import asyncio
import html
import io
import json
import subprocess
from typing import Any

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import HTMLResponse, Response
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from app.core.config import settings
from app.database import session_factory
from app.models.deck import DeckRecord

router = APIRouter(prefix="/api/decks", tags=["exports"])

PptxTheme = tuple[tuple[int, int, int], tuple[int, int, int], tuple[int, int, int]]
PPTX_THEMES: dict[str, PptxTheme] = {
    "ocean": ((242, 248, 252), (22, 50, 74), (42, 121, 165)),
    "sunset": ((253, 244, 232), (75, 42, 31), (198, 92, 44)),
    "forest": ((241, 248, 242), (35, 74, 50), (62, 121, 81)),
}


def _pptx_box(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    fill: tuple[int, int, int] | None = None,
    line: tuple[int, int, int] | None = None,
) -> Any:
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*(fill or (255, 255, 255)))  # type: ignore[no-untyped-call]
    if line:
        shape.line.color.rgb = RGBColor(*line)  # type: ignore[no-untyped-call]
    else:
        shape.line.fill.background()
    return shape


def _pptx_text(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: int,
    color: tuple[int, int, int],
    bold: bool = False,
    font: str = "Aptos",
    align: Any = PP_ALIGN.LEFT,
) -> Any:
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)  # type: ignore[no-untyped-call]
    return box


def _pptx_block_text(block: dict[str, Any]) -> str:
    kind = block["type"]
    if kind in {"heading", "paragraph", "quote"}:
        return str(block.get("text", ""))
    if kind == "stat":
        return f"{block.get('value', '')}\n{block.get('label', '')}\n{block.get('context') or ''}".strip()
    return ""


def _pptx_render_content(
    slide: Any,
    card: dict[str, Any],
    ink: tuple[int, int, int],
    primary: tuple[int, int, int],
    secondary: tuple[int, int, int],
) -> None:
    blocks = card["blocks"]
    layout = card["layout"]
    heading = next((b for b in blocks if b["type"] == "heading"), None)
    paragraphs = [b for b in blocks if b["type"] == "paragraph"]
    bullets = next((b for b in blocks if b["type"] == "bullets"), None)
    columns = next((b for b in blocks if b["type"] == "columns"), None)
    stats = [b for b in blocks if b["type"] == "stat"]
    quotes = [b for b in blocks if b["type"] == "quote"]
    x, y, width = 0.85, 1.15, 11.65

    if layout == "title":
        if paragraphs:
            _pptx_text(slide, paragraphs[0]["text"], x, 0.32, 5.8, 0.32, size=11, color=primary)
        if heading:
            _pptx_text(
                slide, heading["text"], x, 2.55, 10.8, 0.85, size=38, color=ink, font="Georgia"
            )
        for index, paragraph in enumerate(paragraphs[1:2]):
            _pptx_text(
                slide, paragraph["text"], x, 3.55 + index * 0.45, 9.7, 0.4, size=16, color=primary
            )
        return

    if layout == "quote" and quotes:
        _pptx_text(
            slide,
            f"“{quotes[0]['text']}”",
            1.25,
            2.05,
            10.8,
            1.65,
            size=29,
            color=ink,
            font="Georgia",
            align=PP_ALIGN.CENTER,
        )
        if quotes[0].get("attribution"):
            _pptx_text(
                slide,
                quotes[0]["attribution"],
                3.5,
                4.0,
                6.3,
                0.35,
                size=12,
                color=primary,
                align=PP_ALIGN.CENTER,
            )
        return

    if heading:
        _pptx_text(slide, heading["text"], x, y, width, 0.6, size=25, color=ink, bold=True)
        y += 0.85

    if layout in {"three_column", "timeline"} and columns:
        gap = 0.2
        card_width = (width - gap * (len(columns["columns"]) - 1)) / len(columns["columns"])
        for index, column in enumerate(columns["columns"]):
            left = x + index * (card_width + gap)
            _pptx_box(slide, left, 2.05, card_width, 2.15, fill=secondary, line=primary)
            _pptx_text(
                slide,
                column["title"],
                left + 0.2,
                2.3,
                card_width - 0.4,
                0.3,
                size=14,
                color=primary,
                bold=True,
            )
            _pptx_text(
                slide, column["body"], left + 0.2, 2.8, card_width - 0.4, 1.0, size=12, color=ink
            )
        return

    if layout == "two_column" and (bullets or stats):
        if bullets:
            text = "\n".join(f"• {item}" for item in bullets["items"])
            _pptx_text(slide, text, x, y, 6.1, 2.2, size=16, color=ink)
        if stats:
            stat = stats[0]
            _pptx_box(slide, 7.25, 2.0, 2.95, 1.95, fill=secondary, line=primary)
            _pptx_text(
                slide, stat["value"], 7.55, 2.35, 2.35, 0.5, size=29, color=primary, font="Georgia"
            )
            _pptx_text(slide, stat["label"], 7.55, 2.95, 2.35, 0.3, size=12, color=ink, bold=True)
            if stat.get("context"):
                _pptx_text(slide, stat["context"], 7.55, 3.35, 2.35, 0.4, size=9, color=primary)
        return

    if columns:
        gap = 0.2
        card_width = (width - gap * (len(columns["columns"]) - 1)) / len(columns["columns"])
        for index, column in enumerate(columns["columns"]):
            left = x + index * (card_width + gap)
            _pptx_box(slide, left, 2.0, card_width, 1.8, fill=secondary, line=primary)
            _pptx_text(
                slide,
                column["title"],
                left + 0.2,
                2.25,
                card_width - 0.4,
                0.3,
                size=13,
                color=primary,
                bold=True,
            )
            _pptx_text(
                slide, column["body"], left + 0.2, 2.7, card_width - 0.4, 0.75, size=11, color=ink
            )
    elif bullets:
        _pptx_text(
            slide,
            "\n".join(f"• {item}" for item in bullets["items"]),
            x,
            y,
            width,
            2.2,
            size=16,
            color=ink,
        )
    elif paragraphs:
        _pptx_text(
            slide, "\n".join(p["text"] for p in paragraphs), x, y, width, 2.0, size=16, color=ink
        )
    elif stats:
        stat = stats[0]
        _pptx_box(slide, x, y, 3.2, 1.8, fill=secondary, line=primary)
        _pptx_text(
            slide,
            stat["value"],
            x + 0.25,
            y + 0.3,
            2.7,
            0.45,
            size=29,
            color=primary,
            font="Georgia",
        )
        _pptx_text(slide, stat["label"], x + 0.25, y + 0.9, 2.7, 0.3, size=12, color=ink, bold=True)


async def _deck(deck_id: str) -> dict[str, Any]:
    async with session_factory() as session:
        record = await session.get(DeckRecord, deck_id)
    if record is None:
        raise HTTPException(404, "Deck not found")
    return {**record.document, "id": record.id, "title": record.title, "theme": record.theme}


def standalone_html(deck: dict[str, Any]) -> str:
    cards = []
    for index, card in enumerate(deck["cards"], 1):
        blocks = []
        for block in card["blocks"]:
            kind = block["type"]
            if kind == "heading":
                tag = "h1" if block.get("level", 2) == 1 else "h2"
                blocks.append(f"<{tag}>{html.escape(block['text'])}</{tag}>")
            elif kind == "paragraph":
                blocks.append(f"<p>{html.escape(block['text'])}</p>")
            elif kind == "bullets":
                items = "".join(f"<li>{html.escape(item)}</li>" for item in block["items"])
                blocks.append(f"<ul>{items}</ul>")
            elif kind == "quote":
                blocks.append(f"<blockquote>{html.escape(block['text'])}</blockquote>")
            elif kind == "stat":
                blocks.append(
                    f"<p class='stat'><b>{html.escape(block['value'])}</b> "
                    f"{html.escape(block['label'])}</p>"
                )
            else:
                blocks.append(f"<pre>{html.escape(json.dumps(block, ensure_ascii=False))}</pre>")
        cards.append(
            f"<article><div>{''.join(blocks)}</div><footer>"
            f"{html.escape(card['title'])} · {index}/{len(deck['cards'])}</footer></article>"
        )
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(deck["title"])}</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#eee;font-family:Arial,sans-serif;color:#17243b}}
article{{width:297mm;height:210mm;margin:10mm auto;padding:22mm;display:flex;flex-direction:column;justify-content:space-between;background:white;page-break-after:always;box-shadow:0 2px 12px #0002}}
h1{{font-size:42pt}}h2{{font-size:28pt}}p,li{{font-size:18pt;line-height:1.45}}blockquote{{font-size:26pt;border-left:5px solid #5d82bb;padding-left:20px}}footer{{color:#71809b;font-size:11pt;border-top:1px solid #dde3ee;padding-top:10px}}
@media print{{body{{background:white}}article{{margin:0;box-shadow:none}}}}@page{{size:A4 landscape;margin:0}}
</style></head><body>{"".join(cards)}</body></html>"""


@router.get("/{deck_id}/export/html")
async def export_html(deck_id: str) -> Response:
    deck = await _deck(deck_id)
    return HTMLResponse(
        standalone_html(deck),
        headers={"Content-Disposition": f'attachment; filename="{deck["title"]}.html"'},
    )


async def _playwright_bytes(deck_id: str, kind: str, base_url: str) -> bytes:
    try:
        from playwright.async_api import async_playwright
    except ImportError as exc:
        raise HTTPException(
            503, "Playwright is not installed; PDF/PNG export is unavailable."
        ) from exc
    url = f"{base_url.rstrip('/')}/print/{deck_id}"
    async with async_playwright() as playwright:
        browser = None
        for channel in ("chrome", "msedge"):
            try:
                browser = await playwright.chromium.launch(channel=channel, headless=True)
                break
            except Exception:
                continue
        if browser is None:
            browsers_path = settings.resolved_data_dir() / "browsers"
            browsers_path.mkdir(parents=True, exist_ok=True)
            env = {"PLAYWRIGHT_BROWSERS_PATH": str(browsers_path)}
            try:
                await asyncio.to_thread(
                    subprocess.run,
                    ["playwright", "install", "chromium"],
                    env={**__import__("os").environ, **env},
                    check=True,
                    capture_output=True,
                    timeout=300,
                )
                browser = await playwright.chromium.launch(
                    env={"PLAYWRIGHT_BROWSERS_PATH": str(browsers_path)}, headless=True
                )
            except Exception as exc:
                raise HTTPException(
                    503,
                    "No Chrome/Edge found and Chromium download failed. Install a browser "
                    "or Playwright Chromium.",
                ) from exc
        try:
            page = await browser.new_page(
                viewport={"width": 1123, "height": 794}, device_scale_factor=2
            )
            await page.goto(url, wait_until="networkidle", timeout=30_000)
            if kind == "pdf":
                return await page.pdf(
                    format="A4", landscape=True, print_background=True, page_ranges="1-"
                )
            return await page.screenshot(type="png", full_page=True)
        finally:
            await browser.close()


@router.get("/{deck_id}/export/pptx")
async def export_pptx(deck_id: str) -> Response:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise HTTPException(
            503, "python-pptx is not installed; PPTX export is unavailable."
        ) from exc
    deck = await _deck(deck_id)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    background, ink, primary = PPTX_THEMES.get(deck["theme"], PPTX_THEMES["ocean"])
    secondary: tuple[int, int, int] = (
        round((background[0] + primary[0]) / 2),
        round((background[1] + primary[1]) / 2),
        round((background[2] + primary[2]) / 2),
    )
    for card in deck["cards"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*background)  # type: ignore[no-untyped-call]
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor(*primary)  # type: ignore[no-untyped-call]
        accent.line.fill.background()
        _pptx_render_content(slide, card, ink, primary, secondary)
        footer_y = 6.88
        footer = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(footer_y), Inches(11.65), Inches(0.01)
        )
        footer.fill.solid()
        footer.fill.fore_color.rgb = RGBColor(*primary)  # type: ignore[no-untyped-call]
        footer.line.fill.background()
        _pptx_text(slide, card["title"], 0.85, 6.98, 9.6, 0.22, size=9, color=primary)
        _pptx_text(
            slide,
            str(len(presentation.slides)),
            11.8,
            6.98,
            0.7,
            0.22,
            size=9,
            color=primary,
            align=PP_ALIGN.RIGHT,
        )
    output = io.BytesIO()
    presentation.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{deck["title"]}.pptx"'},
    )


@router.get("/{deck_id}/export/{kind}")
async def export_visual(deck_id: str, kind: str, request: Request) -> Response:
    if kind not in {"pdf", "png"}:
        raise HTTPException(404, "Supported formats are pdf and png")
    data = await _playwright_bytes(deck_id, kind, str(request.base_url))
    media = "application/pdf" if kind == "pdf" else "image/png"
    return Response(
        data,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{deck_id}.{kind}"'},
    )
