import io

import pytest
from fastapi import HTTPException, UploadFile
from pptx import Presentation

from app.api.exports import standalone_html
from app.api.imports import import_file, validate_public_url
from app.schemas.generation import Outline


def test_standalone_html_contains_all_cards_and_escaped_text() -> None:
    document = {
        "title": "A <deck>",
        "theme": "ocean",
        "cards": [
            {
                "title": "Intro",
                "blocks": [{"type": "paragraph", "text": "<safe>"}],
            }
        ],
    }
    rendered = standalone_html(document)
    assert "<safe>" not in rendered
    assert "&lt;safe&gt;" in rendered
    assert "Intro" in rendered


@pytest.mark.parametrize(
    "url", ["http://127.0.0.1:8000", "http://localhost", "http://169.254.169.254"]
)
def test_ssrf_guard_blocks_local_and_link_local_urls(url: str) -> None:
    with pytest.raises(HTTPException) as error:
        validate_public_url(url)
    assert error.value.status_code == 400


def test_ssrf_guard_rejects_non_http_schemes() -> None:
    with pytest.raises(HTTPException):
        validate_public_url("file:///etc/passwd")


@pytest.mark.asyncio
async def test_pptx_import_extracts_slide_and_table_text(monkeypatch: pytest.MonkeyPatch) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Quarterly results"
    table = slide.shapes.add_table(2, 2, 0, 0, 4_000_000, 2_000_000).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "$10M"
    buffer = io.BytesIO()
    presentation.save(buffer)
    upload = UploadFile(filename="results.pptx", file=io.BytesIO(buffer.getvalue()))

    captured = ""

    async def fake_outline(text: str) -> Outline:
        nonlocal captured
        captured = text
        return Outline(
            title="Imported presentation",
            items=[{"id": "slide-1", "title": "Slide 1", "summary": text}],
        )

    monkeypatch.setattr("app.api.imports._outline", fake_outline)
    result = await import_file(upload)
    assert result.title == "Imported presentation"
    assert "Quarterly results" in captured
    assert "Revenue | $10M" in captured
